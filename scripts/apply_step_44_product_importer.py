from pathlib import Path

ROOT = Path(r"C:\Projects\sunwise-koyash")


def write_file(relative_path: str, content: str) -> None:
    path = ROOT / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.lstrip(), encoding="utf-8")


def append_once(relative_path: str, line: str) -> None:
    path = ROOT / relative_path
    text = path.read_text(encoding="utf-8") if path.exists() else ""
    if line not in text:
        text = text.rstrip() + "\n" + line + "\n"
        path.write_text(text, encoding="utf-8")


def patch_docker_compose() -> None:
    path = ROOT / "docker-compose.yml"
    text = path.read_text(encoding="utf-8")

    if "./data:/data:ro" in text:
        return

    marker = "      - ./backend:/app"
    if marker not in text:
        raise RuntimeError("Could not find backend volume marker in docker-compose.yml")

    text = text.replace(marker, marker + "\n      - ./data:/data:ro")
    path.write_text(text, encoding="utf-8")


append_once("backend/requirements.txt", "openpyxl>=3.1,<4.0")
patch_docker_compose()

write_file("data/.gitkeep", "\n")
write_file("backend/app/importers/__init__.py", "\n")


write_file("backend/app/importers/product_dataset.py", r'''
import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


COLUMN_ALIASES = {
    "external_id": [
        "external_id",
        "id",
        "product_id",
        "sku",
        "артикул",
        "код",
        "id продукта",
    ],
    "name": [
        "name",
        "product_name",
        "название",
        "название продукта",
        "продукт",
        "средство",
    ],
    "brand": [
        "brand",
        "бренд",
        "марка",
        "производитель",
    ],
    "category": [
        "category",
        "категория",
        "тип продукта",
    ],
    "routine_step": [
        "routine_step",
        "step",
        "шаг",
        "шаг ухода",
        "этап ухода",
    ],
    "price_rub": [
        "price_rub",
        "price",
        "цена",
        "стоимость",
        "цена руб",
        "цена ₽",
    ],
    "budget_segment": [
        "budget_segment",
        "budget",
        "бюджет",
        "ценовой сегмент",
        "сегмент",
    ],
    "skin_types": [
        "skin_types",
        "skin_type",
        "тип кожи",
        "типы кожи",
        "для кожи",
    ],
    "concerns": [
        "concerns",
        "concern",
        "запросы",
        "запрос",
        "проблемы",
        "что волнует",
        "эффект",
        "назначение",
    ],
    "ingredients": [
        "ingredients",
        "ingredient",
        "состав",
        "ингредиенты",
        "ключевые ингредиенты",
        "inci",
    ],
    "ingredient_flags": [
        "ingredient_flags",
        "флаги ингредиентов",
        "свойства состава",
    ],
    "risk_flags": [
        "risk_flags",
        "риски",
        "риск",
        "флаги риска",
        "ограничения",
    ],
    "values": [
        "values",
        "ценности",
        "приоритеты",
        "важно",
    ],
    "texture": [
        "texture",
        "format",
        "формат",
        "текстура",
    ],
    "description": [
        "description",
        "описание",
        "комментарий",
    ],
    "how_to_use": [
        "how_to_use",
        "usage",
        "как использовать",
        "применение",
    ],
    "image_url": [
        "image_url",
        "image",
        "картинка",
        "изображение",
        "фото",
    ],
    "product_url": [
        "product_url",
        "url",
        "ссылка",
        "ссылка на продукт",
    ],
    "is_active": [
        "is_active",
        "active",
        "активен",
        "включен",
        "использовать",
    ],
}


STEP_ALIASES = {
    "cleanser": [
        "cleanser",
        "cleansing",
        "очищение",
        "умывание",
        "гель для умывания",
        "пенка",
        "очищающий гель",
    ],
    "toner": [
        "toner",
        "тоник",
        "тонер",
    ],
    "serum": [
        "serum",
        "сыворотка",
        "актив",
    ],
    "moisturizer": [
        "moisturizer",
        "cream",
        "крем",
        "увлажнение",
        "увлажняющий крем",
        "эмульсия",
    ],
    "spf": [
        "spf",
        "sunscreen",
        "санскрин",
        "солнцезащита",
        "защита от солнца",
    ],
    "exfoliation": [
        "exfoliation",
        "exfoliant",
        "пилинг",
        "эксфолиация",
        "кислоты",
    ],
    "mask": [
        "mask",
        "маска",
    ],
}


SKIN_TYPE_ALIASES = {
    "dry": ["dry", "сухая", "сухой"],
    "oily": ["oily", "жирная", "жирный"],
    "combination": ["combination", "комбинированная", "комбинированный"],
    "normal": ["normal", "нормальная", "нормальный"],
    "sensitive": ["sensitive", "чувствительная", "чувствительный"],
    "all": ["all", "any", "все", "любой", "любая"],
}


CONCERN_ALIASES = {
    "sebum": ["sebum", "жирный блеск", "жирность", "сальность"],
    "acne": ["acne", "акне", "высыпания", "прыщи", "постакне"],
    "dryness": ["dryness", "сухость", "шелушение"],
    "dehydration": ["dehydration", "обезвоженность"],
    "reactivity": ["reactivity", "реактивность", "чувствительность", "раздражение"],
    "redness": ["redness", "покраснение", "краснота"],
    "pigmentation": ["pigmentation", "пигментация", "пятна"],
    "pores": ["pores", "поры", "расширенные поры"],
    "age_changes": ["age_changes", "возрастные изменения", "морщины"],
    "dullness": ["dullness", "тусклость", "серый тон"],
}


VALUE_ALIASES = {
    "comfort": ["comfort", "комфорт"],
    "price": ["price", "цена", "бюджет"],
    "price_quality": ["price_quality", "цена качество", "цена/качество"],
    "quality": ["quality", "качество"],
    "simple_routine": ["simple_routine", "простая рутина", "минимализм"],
    "fast_result": ["fast_result", "быстрый результат"],
    "soft_formula": ["soft_formula", "мягкий состав", "мягкая формула"],
}


FRAGRANCE_MARKERS = {
    "fragrance",
    "parfum",
    "perfume",
    "отдушка",
    "отдушки",
    "ароматизатор",
}

ALCOHOL_MARKERS = {
    "alcohol",
    "ethanol",
    "alcohol denat",
    "denatured alcohol",
    "sd alcohol",
    "спирт",
    "этанол",
}

ESSENTIAL_OIL_MARKERS = {
    "essential oil",
    "essential oils",
    "эфирное масло",
    "эфирные масла",
}


@dataclass(frozen=True)
class ImportErrorItem:
    row_number: int
    reason: str
    row: dict[str, Any]


@dataclass(frozen=True)
class ImportPreview:
    products: list[dict[str, Any]]
    errors: list[ImportErrorItem]


def normalize_header(value: str) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("ё", "е")
    text = re.sub(r"[\s\-_./:;()]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _alias_lookup() -> dict[str, str]:
    result: dict[str, str] = {}

    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            result[normalize_header(alias)] = canonical

    return result


ALIAS_LOOKUP = _alias_lookup()


def _clean_value(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip()

    if text.lower() in {"nan", "none", "null"}:
        return ""

    return text


def _split_list(value: Any) -> list[str]:
    text = _clean_value(value)

    if not text:
        return []

    chunks = re.split(r"[;\n|]+|,\s+", text)
    result = []

    for chunk in chunks:
        item = chunk.strip()
        if item:
            result.append(item)

    return result


def _normalize_by_alias(value: str, aliases: dict[str, list[str]], default: str | None = None) -> str:
    text = normalize_header(value)

    if not text:
        return default or ""

    for canonical, variants in aliases.items():
        normalized_variants = {normalize_header(item) for item in variants}
        if text == canonical or text in normalized_variants:
            return canonical

    return default or text.replace(" ", "_")


def _normalize_list_by_alias(values: list[str], aliases: dict[str, list[str]]) -> list[str]:
    result = []

    for value in values:
        normalized = _normalize_by_alias(value, aliases)
        if normalized and normalized not in result:
            result.append(normalized)

    return result


def _parse_price(value: Any) -> int:
    text = _clean_value(value)

    if not text:
        return 0

    text = text.replace(",", ".")
    match = re.search(r"\d+(?:\.\d+)?", text.replace(" ", ""))

    if not match:
        return 0

    return int(float(match.group(0)))


def _parse_bool(value: Any, default: bool = True) -> bool:
    text = normalize_header(_clean_value(value))

    if not text:
        return default

    if text in {"1", "true", "yes", "y", "да", "активен", "включен", "использовать"}:
        return True

    if text in {"0", "false", "no", "n", "нет", "неактивен", "выключен", "не использовать"}:
        return False

    return default


def _budget_from_price(price: int) -> str:
    if price <= 0:
        return "balanced"

    if price <= 900:
        return "minimum"

    if price <= 2200:
        return "balanced"

    return "premium"


def _slug(value: str) -> str:
    text = normalize_header(value)
    text = re.sub(r"[^a-z0-9а-я]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_")


def _external_id(row: dict[str, Any], brand: str, name: str, row_number: int) -> str:
    existing = _clean_value(row.get("external_id"))

    if existing:
        return _slug(existing) or existing

    basis = f"{brand}-{name}-{row_number}"
    slug = _slug(basis)

    if slug:
        return slug[:80]

    digest = hashlib.sha1(basis.encode("utf-8")).hexdigest()[:12]
    return f"dataset_{digest}"


def _canonical_row(row: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}

    for key, value in row.items():
        normalized_key = normalize_header(str(key))
        canonical = ALIAS_LOOKUP.get(normalized_key)

        if canonical and canonical not in result:
            result[canonical] = value

    return result


def _infer_step(row: dict[str, Any], name: str, category: str) -> str:
    explicit = _clean_value(row.get("routine_step"))

    if explicit:
        return _normalize_by_alias(explicit, STEP_ALIASES, default="other")

    combined = f"{category} {name}".lower()

    for canonical, variants in STEP_ALIASES.items():
        for variant in variants:
            if normalize_header(variant) in normalize_header(combined):
                return canonical

    return "other"


def _infer_risk_flags(ingredients: list[str], existing_flags: list[str]) -> list[str]:
    flags = list(dict.fromkeys(existing_flags))
    joined = " | ".join(item.lower() for item in ingredients)

    if any(marker in joined for marker in FRAGRANCE_MARKERS):
        if "contains_fragrance" not in flags:
            flags.append("contains_fragrance")

    if any(marker in joined for marker in ALCOHOL_MARKERS):
        if "contains_alcohol" not in flags:
            flags.append("contains_alcohol")

    if any(marker in joined for marker in ESSENTIAL_OIL_MARKERS):
        if "contains_essential_oils" not in flags:
            flags.append("contains_essential_oils")

    return flags


def normalize_product_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    canonical = _canonical_row(row)

    name = _clean_value(canonical.get("name"))
    brand = _clean_value(canonical.get("brand"))

    if not name:
        raise ValueError("Missing product name")

    if not brand:
        raise ValueError("Missing product brand")

    category = _clean_value(canonical.get("category"))
    routine_step = _infer_step(canonical, name=name, category=category)
    price = _parse_price(canonical.get("price_rub"))

    ingredients = _split_list(canonical.get("ingredients"))
    ingredient_flags = _normalize_list_by_alias(_split_list(canonical.get("ingredient_flags")), {})
    risk_flags = _infer_risk_flags(
        ingredients=ingredients,
        existing_flags=_normalize_list_by_alias(_split_list(canonical.get("risk_flags")), {}),
    )

    budget_segment = _clean_value(canonical.get("budget_segment"))
    if budget_segment:
        budget_segment = _normalize_by_alias(
            budget_segment,
            {
                "minimum": ["minimum", "low", "низкий", "минимальный", "дешевле"],
                "balanced": ["balanced", "middle", "средний", "баланс"],
                "premium": ["premium", "high", "высокий", "премиум"],
            },
            default="balanced",
        )
    else:
        budget_segment = _budget_from_price(price)

    product = {
        "external_id": _external_id(canonical, brand=brand, name=name, row_number=row_number),
        "name": name,
        "brand": brand,
        "category": category or routine_step,
        "routine_step": routine_step,
        "price_rub": price,
        "currency": "RUB",
        "budget_segment": budget_segment,
        "skin_types": _normalize_list_by_alias(_split_list(canonical.get("skin_types")), SKIN_TYPE_ALIASES),
        "concerns": _normalize_list_by_alias(_split_list(canonical.get("concerns")), CONCERN_ALIASES),
        "ingredients": ingredients,
        "ingredient_flags": ingredient_flags,
        "risk_flags": risk_flags,
        "values": _normalize_list_by_alias(_split_list(canonical.get("values")), VALUE_ALIASES),
        "texture": _clean_value(canonical.get("texture")),
        "description": _clean_value(canonical.get("description")),
        "how_to_use": _clean_value(canonical.get("how_to_use")),
        "image_url": _clean_value(canonical.get("image_url")),
        "product_url": _clean_value(canonical.get("product_url")),
        "is_active": _parse_bool(canonical.get("is_active"), default=True),
        "source": "dataset",
        "source_row": row_number,
    }

    if not product["skin_types"]:
        product["skin_types"] = ["all"]

    if not product["concerns"]:
        product["concerns"] = []

    return product


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    raw_text = path.read_text(encoding="utf-8-sig")
    sample = raw_text[:2048]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")

    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, dialect=dialect)
        for row in reader:
            rows.append(dict(row))

    return rows


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    headers = next(rows_iter, None)

    if not headers:
        return []

    header_values = [str(item or "").strip() for item in headers]
    rows = []

    for values in rows_iter:
        row = {
            header: value
            for header, value in zip(header_values, values, strict=False)
            if header
        }

        if any(_clean_value(value) for value in row.values()):
            rows.append(row)

    return rows


def read_dataset_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported dataset extension: {suffix}")

    if suffix == ".csv":
        return read_csv_rows(path)

    return read_xlsx_rows(path)


def load_products_from_dataset(path: str | Path) -> ImportPreview:
    dataset_path = Path(path)

    if not dataset_path.exists():
        raise FileNotFoundError(f"Dataset file not found: {dataset_path}")

    rows = read_dataset_rows(dataset_path)
    products: list[dict[str, Any]] = []
    errors: list[ImportErrorItem] = []

    seen_external_ids: set[str] = set()

    for index, row in enumerate(rows, start=2):
        try:
            product = normalize_product_row(row, row_number=index)
        except ValueError as exc:
            errors.append(
                ImportErrorItem(
                    row_number=index,
                    reason=str(exc),
                    row=row,
                )
            )
            continue

        external_id = product["external_id"]

        if external_id in seen_external_ids:
            product["external_id"] = f"{external_id}_{index}"

        seen_external_ids.add(product["external_id"])
        products.append(product)

    return ImportPreview(products=products, errors=errors)
''')


write_file("backend/scripts/import_products.py", r'''
import argparse
import asyncio
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from pymongo import AsyncMongoClient

from app.core.config import get_settings
from app.importers.product_dataset import load_products_from_dataset


SAMPLE_PRODUCTS_QUERY = {
    "$or": [
        {"source": "sample"},
        {
            "external_id": {
                "$regex": "sample|cleanser_00|serum_00|moisturizer_00|spf_00",
                "$options": "i",
            }
        },
        {
            "brand": {
                "$regex": "Koyash Sample|Soft Lab|Clear Skin|Sun Basic",
                "$options": "i",
            }
        },
    ]
}


def _json_default(value: Any) -> str:
    return str(value)


async def _delete_count(collection, query: dict[str, Any], dry_run: bool) -> int:
    count = await collection.count_documents(query)

    if dry_run:
        return count

    await collection.delete_many(query)
    return count


async def import_products(args: argparse.Namespace) -> dict[str, Any]:
    settings = get_settings()
    client: AsyncMongoClient[Any] = AsyncMongoClient(
        settings.mongo_uri,
        uuidRepresentation="standard",
    )

    try:
        await client.admin.command("ping")
        db = client[settings.mongo_db_name]
        products_collection = db["products"]

        report: dict[str, Any] = {
            "database": settings.mongo_db_name,
            "dry_run": args.dry_run,
            "file": args.file,
            "sample_products_deleted": 0,
            "dataset_products_deleted": 0,
            "rows_valid": 0,
            "rows_invalid": 0,
            "inserted": 0,
            "updated": 0,
            "errors": [],
        }

        if args.delete_sample_products:
            report["sample_products_deleted"] = await _delete_count(
                products_collection,
                SAMPLE_PRODUCTS_QUERY,
                dry_run=args.dry_run,
            )

        if args.replace_dataset:
            report["dataset_products_deleted"] = await _delete_count(
                products_collection,
                {"source": "dataset"},
                dry_run=args.dry_run,
            )

        if not args.file:
            report["products_total_after"] = await products_collection.count_documents({})
            return report

        preview = load_products_from_dataset(Path(args.file))
        report["rows_valid"] = len(preview.products)
        report["rows_invalid"] = len(preview.errors)
        report["errors"] = [
            {
                "row_number": error.row_number,
                "reason": error.reason,
            }
            for error in preview.errors[:20]
        ]

        if args.dry_run:
            report["preview_products"] = preview.products[:3]
            report["products_total_after"] = await products_collection.count_documents({})
            return report

        now = datetime.now(UTC)

        for product in preview.products:
            document = {
                **product,
                "updated_at": now,
                "imported_at": now,
            }

            result = await products_collection.update_one(
                {"external_id": product["external_id"]},
                {
                    "$set": document,
                    "$setOnInsert": {"created_at": now},
                },
                upsert=True,
            )

            if result.upserted_id is not None:
                report["inserted"] += 1
            else:
                report["updated"] += result.modified_count

        report["products_total_after"] = await products_collection.count_documents({})
        report["dataset_products_after"] = await products_collection.count_documents(
            {"source": "dataset"}
        )
        report["sample_like_products_after"] = await products_collection.count_documents(
            SAMPLE_PRODUCTS_QUERY
        )

        return report
    finally:
        await client.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import KOYASH product dataset into MongoDB.")
    parser.add_argument(
        "--file",
        default=None,
        help="Path to CSV/XLSX file. In Docker use /data/products.xlsx or /data/products.csv.",
    )
    parser.add_argument(
        "--replace-dataset",
        action="store_true",
        help="Delete existing products with source=dataset before import.",
    )
    parser.add_argument(
        "--delete-sample-products",
        action="store_true",
        help="Delete local sample/demo products.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate and show import report without changing MongoDB.",
    )

    return parser.parse_args()


async def main() -> None:
    args = parse_args()
    report = await import_products(args)
    print(json.dumps(report, ensure_ascii=False, indent=2, default=_json_default))


if __name__ == "__main__":
    asyncio.run(main())
''')


write_file("backend/tests/test_product_dataset_importer.py", r'''
from pathlib import Path

from app.importers.product_dataset import load_products_from_dataset, normalize_product_row


def test_product_dataset_mapper_accepts_russian_columns() -> None:
    row = {
        "Название": "Мягкая пенка",
        "Бренд": "Demo Brand",
        "Цена": "750 ₽",
        "Шаг ухода": "Очищение",
        "Типы кожи": "комбинированная; жирная",
        "Запросы": "жирный блеск; высыпания",
        "Состав": "glycerin; parfum",
        "Активен": "да",
    }

    product = normalize_product_row(row, row_number=2)

    assert product["name"] == "Мягкая пенка"
    assert product["brand"] == "Demo Brand"
    assert product["price_rub"] == 750
    assert product["routine_step"] == "cleanser"
    assert product["skin_types"] == ["combination", "oily"]
    assert product["concerns"] == ["sebum", "acne"]
    assert "contains_fragrance" in product["risk_flags"]
    assert product["source"] == "dataset"
    assert product["is_active"] is True


def test_product_dataset_mapper_generates_budget_from_price() -> None:
    row = {
        "name": "Basic SPF",
        "brand": "Demo Brand",
        "price": "2600",
        "routine_step": "spf",
    }

    product = normalize_product_row(row, row_number=3)

    assert product["budget_segment"] == "premium"
    assert product["routine_step"] == "spf"


def test_product_dataset_loader_reads_csv(tmp_path: Path) -> None:
    file_path = tmp_path / "products.csv"
    file_path.write_text(
        "Название;Бренд;Цена;Шаг ухода;Типы кожи;Запросы;Состав\n"
        "Пенка;Demo;700;Очищение;комбинированная;жирный блеск;glycerin\n",
        encoding="utf-8",
    )

    preview = load_products_from_dataset(file_path)

    assert len(preview.products) == 1
    assert preview.errors == []
    assert preview.products[0]["routine_step"] == "cleanser"


def test_product_dataset_loader_collects_invalid_rows(tmp_path: Path) -> None:
    file_path = tmp_path / "products.csv"
    file_path.write_text(
        "Название;Бренд;Цена;Шаг ухода\n"
        ";Demo;700;Очищение\n",
        encoding="utf-8",
    )

    preview = load_products_from_dataset(file_path)

    assert preview.products == []
    assert len(preview.errors) == 1
    assert preview.errors[0].reason == "Missing product name"
''')


write_file("docs/product-dataset-import.md", r'''
# KOYASH product dataset import

This document describes Step 4.4 product dataset import.

The importer supports CSV and XLSX files. Put the real catalog into the root data folder:

data/products.xlsx

or:

data/products.csv

Inside Docker the file is available as:

/data/products.xlsx

or:

/data/products.csv

Main command:

docker compose exec backend python scripts/import_products.py --file /data/products.xlsx --replace-dataset --delete-sample-products

Dry run command:

docker compose exec backend python scripts/import_products.py --file /data/products.xlsx --replace-dataset --delete-sample-products --dry-run

Cleanup sample products only:

docker compose exec backend python scripts/import_products.py --delete-sample-products

Supported column names include Russian and English variants:

external_id, id, sku, артикул, код
name, название, название продукта
brand, бренд
category, категория
routine_step, step, шаг ухода, этап ухода
price_rub, price, цена, стоимость
budget_segment, бюджет, ценовой сегмент
skin_types, тип кожи, типы кожи
concerns, запросы, проблемы, что волнует
ingredients, состав, ингредиенты, ключевые ингредиенты, inci
ingredient_flags, флаги ингредиентов
risk_flags, риски, флаги риска
values, ценности, приоритеты
texture, формат, текстура
description, описание
how_to_use, применение, как использовать
image_url, изображение, фото
product_url, ссылка
is_active, активен, использовать

The importer marks real products with:

source = dataset

Sample/demo products are detected by:

source = sample
or sample-like external_id
or sample demo brands such as Koyash Sample, Soft Lab, Clear Skin, Sun Basic

The importer does not use LLM. It only maps product catalog rows into the backend product schema.
''')


print("Step 4.4 product dataset importer files written successfully.")