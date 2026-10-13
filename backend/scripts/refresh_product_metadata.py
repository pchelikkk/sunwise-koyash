from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from urllib.robotparser import RobotFileParser

import httpx
from openpyxl import load_workbook

from app.services.catalog_metadata import (
    canonical_product_url,
    extract_catalog_metadata,
    is_supported_product_url,
)

USER_AGENT = "KOYASHCatalogUpdater/1.0"
MAX_HTML_BYTES = 5_000_000


class RobotsPolicy:
    def __init__(self, client: httpx.Client, *, ignore: bool = False) -> None:
        self.client = client
        self.ignore = ignore
        self.cache: dict[str, tuple[bool, RobotFileParser | None]] = {}

    def allows(self, url: str) -> tuple[bool, str | None]:
        if self.ignore:
            return True, "robots_check_explicitly_disabled"
        parsed = urlparse(url)
        origin = f"{parsed.scheme}://{parsed.netloc}"
        if origin not in self.cache:
            robots_url = f"{origin}/robots.txt"
            try:
                response = self.client.get(robots_url)
                if response.status_code == 404:
                    self.cache[origin] = (True, None)
                elif response.status_code >= 400:
                    self.cache[origin] = (False, None)
                else:
                    parser = RobotFileParser()
                    parser.set_url(robots_url)
                    parser.parse(response.text.splitlines())
                    self.cache[origin] = (True, parser)
            except httpx.HTTPError:
                self.cache[origin] = (False, None)
        robots_available, parser = self.cache[origin]
        if not robots_available:
            return False, "robots_unavailable"
        if parser is not None and not parser.can_fetch(USER_AGENT, url):
            return False, "robots_disallow"
        return True, None


def _column_map(worksheet: Any) -> dict[str, int]:
    return {
        str(cell.value or "").strip().lower(): cell.column
        for cell in worksheet[1]
        if str(cell.value or "").strip()
    }


def _find_column(columns: dict[str, int], *aliases: str) -> int | None:
    return next((columns[alias] for alias in aliases if alias in columns), None)


def _ensure_column(worksheet: Any, columns: dict[str, int], name: str) -> int:
    existing = columns.get(name)
    if existing is not None:
        return existing
    column = worksheet.max_column + 1
    worksheet.cell(row=1, column=column, value=name)
    columns[name] = column
    return column


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _materialize_derived_fields(
    worksheet: Any,
    columns: dict[str, int],
    *,
    apply: bool,
) -> int:
    """Replace volatile Excel formulas with portable, current values.

    openpyxl preserves formulas but cannot recalculate their cached values. The
    backend intentionally reads workbooks with ``data_only=True``, so saving a
    formula workbook without Excel would otherwise make its derived fields look
    empty. Materializing these two documented dataset formulas keeps imports
    deterministic in local development and CI.
    """

    price_column = _find_column(columns, "price_rub", "price")
    segment_column = _find_column(columns, "segment", "budget_segment", "budget")
    volume_column = _find_column(columns, "volume_ml", "volume")
    price_per_ml_column = _find_column(columns, "price_per_ml")
    if price_column is None:
        return 0

    changed = 0
    for row in range(2, worksheet.max_row + 1):
        price = _number(worksheet.cell(row=row, column=price_column).value)
        if price is None or price <= 0:
            continue

        if segment_column is not None:
            segment = "бюджет" if price < 1000 else "мидл" if price <= 3000 else "люкс"
            cell = worksheet.cell(row=row, column=segment_column)
            if cell.value != segment:
                changed += 1
                if apply:
                    cell.value = segment

        if volume_column is not None and price_per_ml_column is not None:
            volume = _number(worksheet.cell(row=row, column=volume_column).value)
            if volume is None or volume <= 0:
                continue
            price_per_ml = price / volume
            cell = worksheet.cell(row=row, column=price_per_ml_column)
            current = _number(cell.value)
            if current is None or abs(current - price_per_ml) > 1e-9:
                changed += 1
                if apply:
                    cell.value = price_per_ml

    return changed


def _save_atomic(workbook: Any, dataset_path: Path) -> None:
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{dataset_path.stem}-",
        suffix=dataset_path.suffix,
        dir=dataset_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, dataset_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _fetch_metadata(
    client: httpx.Client,
    url: str,
    current_price: int | None,
) -> tuple[Any | None, str | None]:
    try:
        response = client.get(url)
        response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        return None, f"http_{exc.response.status_code}"
    except httpx.HTTPError as exc:
        return None, f"network_error:{type(exc).__name__}"

    if len(response.content) > MAX_HTML_BYTES:
        return None, "response_too_large"
    content_type = response.headers.get("content-type", "").lower()
    if "html" not in content_type:
        return None, "not_html"
    final_url = canonical_product_url(str(response.url))
    if final_url is None or not is_supported_product_url(final_url):
        return None, "unsupported_redirect"
    return (
        extract_catalog_metadata(
            response.text,
            page_url=final_url,
            current_price_rub=current_price,
        ),
        None,
    )


def refresh(args: argparse.Namespace) -> dict[str, Any]:
    dataset_path = Path(args.dataset).resolve()
    if not dataset_path.exists():
        raise FileNotFoundError(f"Dataset not found: {dataset_path}")
    if dataset_path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx datasets are supported")

    workbook = load_workbook(dataset_path)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {args.sheet}")
    worksheet = workbook[args.sheet]
    columns = _column_map(worksheet)
    id_column = _find_column(columns, "external_id", "product_id", "id")
    url_column = _find_column(columns, "product_url", "link", "url")
    price_column = _find_column(columns, "price_rub", "price")
    image_column = _find_column(columns, "image_url", "image", "photo_url")
    if id_column is None or price_column is None:
        raise ValueError("Dataset must contain product id and price columns")
    if not args.derived_only and url_column is None:
        raise ValueError("Dataset must contain a product URL column")

    if args.apply and not args.derived_only:
        image_column = image_column or _ensure_column(worksheet, columns, "image_url")
        checked_column = _ensure_column(worksheet, columns, "metadata_checked_at")
        status_column = _ensure_column(worksheet, columns, "metadata_status")
    else:
        checked_column = None
        status_column = None

    only_ids = set(args.only_id or [])
    checked_at = datetime.now(UTC).isoformat()
    report: dict[str, Any] = {
        "ok": True,
        "mode": "apply" if args.apply else "dry_run",
        "dataset": str(dataset_path),
        "sheet": args.sheet,
        "checked_at": checked_at,
        "summary": {
            "eligible": 0,
            "requested": 0,
            "metadata_found": 0,
            "price_changes": 0,
            "image_changes": 0,
            "derived_cells_updated": 0,
            "applied_rows": 0,
            "skipped": 0,
            "errors": 0,
        },
        "items": [],
    }

    if args.derived_only:
        report["summary"]["derived_cells_updated"] = _materialize_derived_fields(
            worksheet,
            columns,
            apply=args.apply,
        )
        if args.apply:
            _save_atomic(workbook, dataset_path)
        return report

    headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.5",
    }
    last_request_at: float | None = None
    with httpx.Client(
        headers=headers,
        follow_redirects=True,
        timeout=httpx.Timeout(args.timeout_seconds),
    ) as client:
        robots = RobotsPolicy(client, ignore=args.ignore_robots)
        for row in range(2, worksheet.max_row + 1):
            product_id = str(worksheet.cell(row=row, column=id_column).value or "").strip()
            if not product_id or (only_ids and product_id not in only_ids):
                continue
            raw_url = worksheet.cell(row=row, column=url_column).value
            url = canonical_product_url(raw_url)
            if not url:
                report["summary"]["skipped"] += 1
                report["items"].append(
                    {"product_id": product_id, "status": "missing_or_invalid_url"}
                )
                continue
            if not is_supported_product_url(url):
                report["summary"]["skipped"] += 1
                report["items"].append(
                    {"product_id": product_id, "status": "unsupported_domain", "url": url}
                )
                continue

            report["summary"]["eligible"] += 1
            if args.limit and report["summary"]["requested"] >= args.limit:
                break
            allowed, robots_note = robots.allows(url)
            if not allowed:
                report["summary"]["skipped"] += 1
                report["items"].append(
                    {"product_id": product_id, "status": robots_note, "url": url}
                )
                continue

            if last_request_at is not None:
                remaining = args.delay_seconds - (time.monotonic() - last_request_at)
                if remaining > 0:
                    time.sleep(remaining)
            current_price_value = worksheet.cell(row=row, column=price_column).value
            try:
                current_price = int(round(float(current_price_value)))
            except (TypeError, ValueError):
                current_price = None

            report["summary"]["requested"] += 1
            metadata, fetch_error = _fetch_metadata(client, url, current_price)
            last_request_at = time.monotonic()
            item: dict[str, Any] = {
                "product_id": product_id,
                "url": url,
                "current_price_rub": current_price,
                "robots_note": robots_note,
            }
            if fetch_error:
                item["status"] = fetch_error
                report["summary"]["errors"] += 1
                report["items"].append(item)
                continue

            item.update(
                {
                    "fetched_price_rub": metadata.price_rub,
                    "fetched_image_url": metadata.image_url,
                    "price_candidates": metadata.price_candidates,
                    "warnings": metadata.warnings,
                }
            )
            if metadata.price_rub is None and metadata.image_url is None:
                item["status"] = "metadata_not_found"
                report["summary"]["skipped"] += 1
                report["items"].append(item)
                continue

            report["summary"]["metadata_found"] += 1
            current_image = (
                str(worksheet.cell(row=row, column=image_column).value or "").strip()
                if image_column
                else ""
            )
            price_changed = bool(
                metadata.price_rub is not None and metadata.price_rub != current_price
            )
            price_guarded = False
            if price_changed and current_price and current_price > 0:
                change_percent = round(
                    abs(metadata.price_rub - current_price) / current_price * 100,
                    2,
                )
                item["price_change_percent"] = change_percent
                price_guarded = change_percent > args.max_price_change_percent
            image_changed = bool(
                metadata.image_url
                and (not current_image or args.replace_existing_images)
                and metadata.image_url != current_image
            )
            if price_changed and not price_guarded:
                report["summary"]["price_changes"] += 1
            if image_changed:
                report["summary"]["image_changes"] += 1

            applied_fields: list[str] = []
            if args.apply:
                if price_changed and not price_guarded:
                    worksheet.cell(row=row, column=price_column, value=metadata.price_rub)
                    applied_fields.append("price_rub")
                if image_changed and image_column:
                    worksheet.cell(row=row, column=image_column, value=metadata.image_url)
                    applied_fields.append("image_url")
                if checked_column and status_column:
                    worksheet.cell(row=row, column=checked_column, value=checked_at)
                    worksheet.cell(
                        row=row,
                        column=status_column,
                        value="price_guard" if price_guarded else "checked",
                    )
                if applied_fields:
                    report["summary"]["applied_rows"] += 1

            item["status"] = "price_change_guarded" if price_guarded else "ok"
            item["applied_fields"] = applied_fields
            report["items"].append(item)

    if args.apply:
        report["summary"]["derived_cells_updated"] = _materialize_derived_fields(
            worksheet,
            columns,
            apply=True,
        )
        _save_atomic(workbook, dataset_path)
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Refresh product prices and images from structured metadata on approved "
            "merchant pages. Dry-run is the default."
        )
    )
    parser.add_argument("--dataset", default="/data/products.xlsx")
    parser.add_argument("--sheet", default="Products")
    parser.add_argument("--only-id", action="append")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--delay-seconds", type=float, default=1.0)
    parser.add_argument("--timeout-seconds", type=float, default=20.0)
    parser.add_argument("--max-price-change-percent", type=float, default=35.0)
    parser.add_argument("--replace-existing-images", action="store_true")
    parser.add_argument("--ignore-robots", action="store_true")
    parser.add_argument(
        "--derived-only",
        action="store_true",
        help=(
            "Only preview or materialize the dataset's segment and price-per-ml "
            "fields; do not make network requests."
        ),
    )
    parser.add_argument("--report")
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Write accepted changes atomically. Without this flag nothing is changed.",
    )
    args = parser.parse_args()
    if args.limit < 0 or args.delay_seconds < 0 or args.timeout_seconds <= 0:
        parser.error("limit and delays must be non-negative; timeout must be positive")
    return args


def main() -> None:
    args = parse_args()
    try:
        report = refresh(args)
    except (FileNotFoundError, ValueError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2))
        raise SystemExit(2) from exc
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    print(rendered)
    if args.report:
        Path(args.report).write_text(rendered + "\n", encoding="utf-8")
    if report["summary"]["errors"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
