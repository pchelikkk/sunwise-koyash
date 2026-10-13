from pathlib import Path

ROOT = Path(r"C:\Projects\sunwise-koyash")


def write_file(relative_path: str, content: str) -> None:
    path = ROOT / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.lstrip(), encoding="utf-8")


write_file("backend/app/importers/product_dataset.py", r'''
import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


COLUMN_ALIASES = {
    "external_id": [
        "external_id",
        "id",
        "product_id",
        "sku",
        "артикул",
        "код",
        "id продукта",
    ],
    "name": [
        "name",
        "product_name",
        "название",
        "название продукта",
        "продукт",
        "средство",
    ],
    "brand": [
        "brand",
        "бренд",
        "марка",
        "производитель",
    ],
    "category": [
        "category",
        "категория",
        "тип продукта",
    ],
    "routine_step": [
        "routine_step",
        "step",
        "шаг",
        "шаг ухода",
        "этап ухода",
    ],
    "price_rub": [
        "price_rub",
        "price",
        "цена",
        "стоимость",
        "цена руб",
        "цена ₽",
    ],
    "budget_segment": [
        "budget_segment",
        "budget",
        "segment",
        "сегмент",
        "бюджет",
        "ценовой сегмент",
    ],
    "skin_types": [
        "skin_types",
        "skin_type",
        "skintype",
        "тип кожи",
        "типы кожи",
        "для кожи",
    ],
    "concerns": [
        "concerns",
        "concern",
        "запросы",
        "запрос",
        "проблемы",
        "что волнует",
        "эффект",
        "назначение",
    ],
    "ingredients": [
        "ingredients",
        "ingredient",
        "ingredients_raw",
        "состав",
        "ингредиенты",
        "ключевые ингредиенты",
        "inci",
    ],
    "key_ingredients": [
        "main_actives",
        "key_ingredients",
        "active_ingredients",
        "активы",
        "основные активы",
        "ключевые активы",
    ],
    "functional_category": [
        "functional_category",
        "functional category",
        "функциональная категория",
        "функции",
        "эффекты",
    ],
    "ingredient_flags": [
        "ingredient_flags",
        "флаги ингредиентов",
        "свойства состава",
    ],
    "risk_flags": [
        "risk_flags",
        "риски",
        "риск",
        "флаги риска",
        "ограничения",
    ],
    "allergens": [
        "allergens",
        "аллергены",
        "раздражители",
    ],
    "values": [
        "values",
        "ценности",
        "приоритеты",
        "важно",
    ],
    "texture": [
        "texture",
        "format",
        "формат",
        "текстура",
    ],
    "description": [
        "description",
        "описание",
        "комментарий",
    ],
    "how_to_use": [
        "how_to_use",
        "usage",
        "как использовать",
        "применение",
    ],
    "image_url": [
        "image_url",
        "image",
        "картинка",
        "изображение",
        "фото",
    ],
    "product_url": [
        "product_url",
        "url",
        "link",
        "ссылка",
        "ссылка на продукт",
    ],
    "is_active": [
        "is_active",
        "active",
        "активен",
        "включен",
        "использовать",
    ],
    "product_status": [
        "status",
        "статус",
    ],
    "ph": [
        "ph",
        "pH",
        "кислотность",
    ],
    "vegan": [
        "vegan",
        "веган",
    ],
    "cruelty_free": [
        "cruelty_free",
        "cruelty free",
        "не тестируется на животных",
    ],
    "volume_ml": [
        "volume_ml",
        "volume",
        "объем",
        "объём",
        "мл",
    ],
    "price_per_ml": [
        "price_per_ml",
        "цена за мл",
    ],
    "issues": [
        "issues",
        "заметки",
        "проблемы состава",
        "комментарии по составу",
    ],
}


STEP_ALIASES = {
    "cleanser": [
        "cleanser",
        "cleansing",
        "очищение",
        "очищающий",
        "умывание",
        "гель для умывания",
        "пенка",
        "мицеллярная вода",
        "мицеллярная пенка",
    ],
    "toner": [
        "toner",
        "тоник",
        "тонер",
        "эссенция",
    ],
    "serum": [
        "serum",
        "сыворотка",
        "актив",
    ],
    "moisturizer": [
        "moisturizer",
        "cream",
        "крем",
        "гель-крем",
        "увлажняющий крем",
        "эмульсия",
        "дневной",
        "ночной",
    ],
    "spf": [
        "spf",
        "sunscreen",
        "санскрин",
        "солнцезащита",
        "солнцезащитный",
        "защита от солнца",
    ],
    "exfoliation": [
        "exfoliation",
        "exfoliant",
        "пилинг",
        "эксфолиация",
        "кислоты",
        "aha",
        "bha",
        "pha",
    ],
    "mask": [
        "mask",
        "маска",
    ],
}


SKIN_TYPE_ALIASES = {
    "dry": ["dry", "сухая", "сухой", "сухость"],
    "oily": ["oily", "жирная", "жирный", "себум", "сальность"],
    "combination": ["combination", "комбинированная", "комбинированный", "t зона", "t-зона"],
    "normal": ["normal", "нормальная", "нормальный"],
    "sensitive": ["sensitive", "чувствительная", "чувствительный", "раздражение"],
    "all": ["all", "any", "все", "любой", "любая"],
}


CONCERN_ALIASES = {
    "sebum": ["sebum", "себум", "жирный блеск", "жирность", "сальность", "матирование", "себорегуляция", "себоконтроль"],
    "acne": ["acne", "акне", "высыпания", "прыщи", "постакне", "черные точки", "чёрные точки", "антибактериальное"],
    "dryness": ["dryness", "сухость", "шелушение", "питание"],
    "dehydration": ["dehydration", "обезвоженность", "гидратация", "увлажнение", "увлажняющий"],
    "reactivity": ["reactivity", "реактивность", "чувствительность", "раздражение", "успокаивание", "воспаление", "противовоспаление"],
    "redness": ["redness", "покраснение", "краснота"],
    "pigmentation": ["pigmentation", "пигментация", "осветление", "пятна"],
    "pores": ["pores", "поры", "расширенные поры", "черные точки", "чёрные точки"],
    "age_changes": ["age_changes", "возрастные изменения", "морщины", "anti age", "анти эйджинг", "упругость", "антивозраст"],
    "dullness": ["dullness", "тусклость", "серый тон", "сияние"],
}


VALUE_ALIASES = {
    "comfort": ["comfort", "комфорт"],
    "price": ["price", "цена", "бюджет"],
    "price_quality": ["price_quality", "цена качество", "цена/качество", "баланс"],
    "quality": ["quality", "качество"],
    "simple_routine": ["simple_routine", "простая рутина", "минимализм", "минималистичный"],
    "fast_result": ["fast_result", "быстрый результат", "результативность"],
    "soft_formula": ["soft_formula", "мягкий состав", "мягкая формула", "деликатное"],
    "vegan": ["vegan", "веган"],
    "cruelty_free": ["cruelty_free", "не тестируется"],
}


FRAGRANCE_MARKERS = {
    "fragrance",
    "parfum",
    "perfume",
    "отдушка",
    "отдушки",
    "ароматизатор",
}

DRYING_ALCOHOL_MARKERS = {
    "alcohol denat",
    "denatured alcohol",
    "sd alcohol",
    "ethanol",
    "ethyl alcohol",
    "этанол",
    "денатурированный спирт",
}

ESSENTIAL_OIL_MARKERS = {
    "essential oil",
    "essential oils",
    "эфирное масло",
    "эфирные масла",
    "bergamot",
    "grapefruit peel oil",
    "geranium",
    "sage oil",
    "patchouli",
    "tea tree",
    "melaleuca alternifolia",
}


@dataclass(frozen=True)
class ImportErrorItem:
    row_number: int
    reason: str
    row: dict[str, Any]


@dataclass(frozen=True)
class ImportPreview:
    products: list[dict[str, Any]]
    errors: list[ImportErrorItem]


def normalize_header(value: str) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("ё", "е")
    text = re.sub(r"[\s\-_./:;()]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _alias_lookup() -> dict[str, str]:
    result: dict[str, str] = {}

    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            result[normalize_header(alias)] = canonical

    return result


ALIAS_LOOKUP = _alias_lookup()


def _clean_value(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip()

    if text.lower() in {"nan", "none", "null"}:
        return ""

    return text


def _split_list(value: Any) -> list[str]:
    text = _clean_value(value)

    if not text:
        return []

    chunks = re.split(r"[;\n|]+|,\s+", text)
    result = []

    for chunk in chunks:
        item = chunk.strip()
        if item:
            result.append(item)

    return result


def _normalize_by_alias(value: str, aliases: dict[str, list[str]], default: str | None = None) -> str:
    text = normalize_header(value)

    if not text:
        return default or ""

    for canonical, variants in aliases.items():
        normalized_variants = {normalize_header(item) for item in variants}

        if text == canonical or text in normalized_variants:
            return canonical

    return default or text.replace(" ", "_")


def _normalize_list_by_alias(values: list[str], aliases: dict[str, list[str]]) -> list[str]:
    result = []

    for value in values:
        normalized = _normalize_by_alias(value, aliases)

        if normalized and normalized not in result:
            result.append(normalized)

    return result


def _parse_price(value: Any) -> int:
    text = _clean_value(value)

    if not text:
        return 0

    text = text.replace(",", ".")
    match = re.search(r"\d+(?:\.\d+)?", text.replace(" ", ""))

    if not match:
        return 0

    return int(float(match.group(0)))


def _parse_float(value: Any) -> float | None:
    text = _clean_value(value)

    if not text:
        return None

    text = text.replace(",", ".")
    match = re.search(r"\d+(?:\.\d+)?", text.replace(" ", ""))

    if not match:
        return None

    return float(match.group(0))


def _parse_bool(value: Any, default: bool = True) -> bool:
    if isinstance(value, bool):
        return value

    text = normalize_header(_clean_value(value))

    if not text:
        return default

    if text in {"1", "true", "yes", "y", "да", "активен", "включен", "использовать"}:
        return True

    if text in {"0", "false", "no", "n", "нет", "неактивен", "выключен", "не использовать"}:
        return False

    return default


def _budget_from_price(price: int) -> str:
    if price <= 0:
        return "balanced"

    if price <= 900:
        return "minimum"

    if price <= 2200:
        return "balanced"

    return "premium"


def _budget_from_segment(value: str, price: int) -> str:
    text = normalize_header(value)

    if text in {"бюджет", "budget", "low", "minimum", "низкий", "минимальный"}:
        return "minimum"

    if text in {"мидл", "middle", "medium", "balanced", "средний", "баланс"}:
        return "balanced"

    if text in {"люкс", "lux", "premium", "high", "премиум", "высокий"}:
        return "premium"

    return _budget_from_price(price)


def _slug(value: str) -> str:
    text = normalize_header(value)
    text = re.sub(r"[^a-z0-9а-я]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_")


def _external_id(row: dict[str, Any], brand: str, name: str, row_number: int) -> str:
    existing = _clean_value(row.get("external_id"))

    if existing:
        return _slug(existing) or existing

    basis = f"{brand}-{name}-{row_number}"
    slug = _slug(basis)

    if slug:
        return slug[:80]

    digest = hashlib.sha1(basis.encode("utf-8")).hexdigest()[:12]
    return f"dataset_{digest}"


def _canonical_row(row: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}

    for key, value in row.items():
        normalized_key = normalize_header(str(key))
        canonical = ALIAS_LOOKUP.get(normalized_key)

        if canonical and canonical not in result:
            result[canonical] = value

    return result


def _infer_step(row: dict[str, Any], name: str, category: str, functional_category: str) -> str:
    explicit = _clean_value(row.get("routine_step"))

    if explicit:
        return _normalize_by_alias(explicit, STEP_ALIASES, default="other")

    combined = normalize_header(f"{category} {name} {functional_category}")

    if "spf" in combined or "солнцезащит" in combined:
        return "spf"

    if "пилинг" in combined or "эксфоли" in combined:
        return "exfoliation"

    if "сыворот" in combined:
        return "serum"

    if "маска" in combined:
        return "mask"

    if "крем" in combined or "увлажн" in combined:
        return "moisturizer"

    if "тоник" in combined or "тонер" in combined or "эссенц" in combined:
        return "toner"

    if "очищ" in combined or "пенка" in combined or "умыван" in combined or "мицелляр" in combined:
        return "cleanser"

    for canonical, variants in STEP_ALIASES.items():
        for variant in variants:
            if normalize_header(variant) in combined:
                return canonical

    return "other"


def _infer_skin_types(text: str) -> list[str]:
    normalized = normalize_header(text)
    result: list[str] = []

    for canonical, aliases in SKIN_TYPE_ALIASES.items():
        for alias in aliases:
            if normalize_header(alias) in normalized and canonical not in result:
                result.append(canonical)

    return result


def _infer_concerns(text: str) -> list[str]:
    normalized = normalize_header(text)
    result: list[str] = []

    for canonical, aliases in CONCERN_ALIASES.items():
        for alias in aliases:
            if normalize_header(alias) in normalized and canonical not in result:
                result.append(canonical)

    return result


def _yes_value(value: Any) -> bool:
    return _parse_bool(value, default=False)


def _infer_values(row: dict[str, Any], functional_category: str, issues: str) -> list[str]:
    values = _normalize_list_by_alias(_split_list(row.get("values")), VALUE_ALIASES)

    if _yes_value(row.get("vegan")) and "vegan" not in values:
        values.append("vegan")

    if _yes_value(row.get("cruelty_free")) and "cruelty_free" not in values:
        values.append("cruelty_free")

    combined = normalize_header(f"{functional_category} {issues}")

    if "мягк" in combined or "деликат" in combined:
        values.append("soft_formula")

    if "баланс" in combined:
        values.append("price_quality")

    return list(dict.fromkeys(values))


def _infer_risk_flags(
    *,
    ingredients: list[str],
    allergens: str,
    issues: str,
    existing_flags: list[str],
) -> list[str]:
    flags = list(dict.fromkeys(existing_flags))
    joined = " | ".join([*ingredients, allergens, issues]).lower()

    if any(marker in joined for marker in FRAGRANCE_MARKERS):
        flags.append("contains_fragrance")

    if any(marker in joined for marker in DRYING_ALCOHOL_MARKERS):
        flags.append("contains_alcohol")

    if any(marker in joined for marker in ESSENTIAL_OIL_MARKERS):
        flags.append("contains_essential_oils")

    if "phototoxic" in joined or "фототокс" in joined:
        flags.append("phototoxic_components")

    if "sensitive" in joined or "чувствитель" in joined:
        flags.append("sensitive_attention")

    return list(dict.fromkeys(flags))


def normalize_product_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    canonical = _canonical_row(row)

    name = _clean_value(canonical.get("name"))
    brand = _clean_value(canonical.get("brand"))

    if not name:
        raise ValueError("Missing product name")

    if not brand:
        raise ValueError("Missing product brand")

    category = _clean_value(canonical.get("category"))
    functional_category = _clean_value(canonical.get("functional_category"))
    key_ingredients_text = _clean_value(canonical.get("key_ingredients"))
    allergens = _clean_value(canonical.get("allergens"))
    issues = _clean_value(canonical.get("issues"))

    routine_step = _infer_step(
        canonical,
        name=name,
        category=category,
        functional_category=functional_category,
    )

    price = _parse_price(canonical.get("price_rub"))
    ingredients = _split_list(canonical.get("ingredients"))
    key_ingredients = _split_list(key_ingredients_text)

    risk_flags = _infer_risk_flags(
        ingredients=ingredients,
        allergens=allergens,
        issues=issues,
        existing_flags=_normalize_list_by_alias(_split_list(canonical.get("risk_flags")), {}),
    )

    budget_segment = _budget_from_segment(
        _clean_value(canonical.get("budget_segment")),
        price=price,
    )

    skin_types = _normalize_list_by_alias(
        _split_list(canonical.get("skin_types")),
        SKIN_TYPE_ALIASES,
    )

    if not skin_types:
        skin_types = _infer_skin_types(f"{name} {category} {functional_category} {issues}")

    if not skin_types:
        skin_types = ["all"]

    concerns = _normalize_list_by_alias(
        _split_list(canonical.get("concerns")),
        CONCERN_ALIASES,
    )

    if not concerns:
        concerns = _infer_concerns(
            f"{name} {category} {functional_category} {key_ingredients_text} {issues}"
        )

    product = {
        "external_id": _external_id(canonical, brand=brand, name=name, row_number=row_number),
        "name": name,
        "brand": brand,
        "category": category or routine_step,
        "routine_step": routine_step,
        "price_rub": price,
        "currency": "RUB",
        "budget_segment": budget_segment,
        "skin_types": skin_types,
        "concerns": concerns,
        "ingredients": ingredients,
        "key_ingredients": key_ingredients,
        "ingredient_highlights": key_ingredients_text,
        "functional_category": functional_category,
        "ingredient_flags": _normalize_list_by_alias(_split_list(canonical.get("ingredient_flags")), {}),
        "risk_flags": risk_flags,
        "allergens": allergens,
        "values": _infer_values(canonical, functional_category=functional_category, issues=issues),
        "texture": _clean_value(canonical.get("texture")),
        "description": _clean_value(canonical.get("description")) or functional_category,
        "how_to_use": _clean_value(canonical.get("how_to_use")),
        "image_url": _clean_value(canonical.get("image_url")),
        "product_url": _clean_value(canonical.get("product_url")),
        "ph": _clean_value(canonical.get("ph")),
        "vegan": _parse_bool(canonical.get("vegan"), default=False),
        "cruelty_free": _parse_bool(canonical.get("cruelty_free"), default=False),
        "volume_ml": _parse_float(canonical.get("volume_ml")),
        "price_per_ml": _parse_float(canonical.get("price_per_ml")),
        "issues": issues,
        "product_status": canonical.get("product_status"),
        "is_active": _parse_bool(canonical.get("is_active"), default=True),
        "source": "dataset",
        "source_row": row_number,
    }

    return product


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    raw_text = path.read_text(encoding="utf-8-sig")
    sample = raw_text[:2048]

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = []

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, dialect=dialect)

        for row in reader:
            rows.append(dict(row))

    return rows


def _header_score(headers: list[Any], sheet_name: str) -> int:
    canonical_headers = set()

    for header in headers:
        normalized = normalize_header(str(header or ""))
        canonical = ALIAS_LOOKUP.get(normalized)

        if canonical:
            canonical_headers.add(canonical)

    score = len(canonical_headers)

    if {"external_id", "name", "brand"}.issubset(canonical_headers):
        score += 20

    if normalize_header(sheet_name) in {"products", "product", "товары", "продукты", "catalog", "каталог"}:
        score += 10

    return score


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)

    best_sheet = None
    best_header_index = 0
    best_headers = None
    best_score = -1

    for sheet in workbook.worksheets:
        preview_rows = []

        for row in sheet.iter_rows(values_only=True):
            preview_rows.append(row)

            if len(preview_rows) >= 10:
                break

        for index, row in enumerate(preview_rows):
            headers = [str(item or "").strip() for item in row]
            score = _header_score(headers, sheet.title)

            if score > best_score:
                best_score = score
                best_sheet = sheet
                best_header_index = index
                best_headers = headers

    if best_sheet is None or best_headers is None:
        return []

    rows = []

    for row_index, values in enumerate(best_sheet.iter_rows(values_only=True)):
        if row_index <= best_header_index:
            continue

        row = {
            header: value
            for header, value in zip(best_headers, values, strict=False)
            if header
        }

        if any(_clean_value(value) for value in row.values()):
            rows.append(row)

    return rows


def read_dataset_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported dataset extension: {suffix}")

    if suffix == ".csv":
        return read_csv_rows(path)

    return read_xlsx_rows(path)


def load_products_from_dataset(path: str | Path) -> ImportPreview:
    dataset_path = Path(path)

    if not dataset_path.exists():
        raise FileNotFoundError(f"Dataset file not found: {dataset_path}")

    rows = read_dataset_rows(dataset_path)
    products: list[dict[str, Any]] = []
    errors: list[ImportErrorItem] = []

    seen_external_ids: set[str] = set()

    for index, row in enumerate(rows, start=2):
        try:
            product = normalize_product_row(row, row_number=index)
        except ValueError as exc:
            errors.append(
                ImportErrorItem(
                    row_number=index,
                    reason=str(exc),
                    row=row,
                )
            )
            continue

        external_id = product["external_id"]

        if external_id in seen_external_ids:
            product["external_id"] = f"{external_id}_{index}"

        seen_external_ids.add(product["external_id"])
        products.append(product)

    return ImportPreview(products=products, errors=errors)
''')


write_file("backend/tests/test_product_dataset_importer.py", r'''
from pathlib import Path

from openpyxl import Workbook

from app.importers.product_dataset import load_products_from_dataset, normalize_product_row


def test_product_dataset_mapper_accepts_actual_products_columns() -> None:
    row = {
        "product_id": "product_001",
        "name": "Мягкая пенка",
        "brand": "Demo Brand",
        "price_rub": 743,
        "segment": "бюджет",
        "category": "Очищающий гель/пенка",
        "link": "https://example.com/product",
        "ingredients_raw": "Aqua, Glycerin, Fragrance",
        "main_actives": "Glycerin (увлажнитель)",
        "functional_category": "Мягкое очищение, гидратация",
        "allergens": "Fragrance",
        "vegan": "Да",
        "cruelty_free": "Да",
        "format": "Пенка",
        "volume_ml": 75,
        "price_per_ml": 9.9,
        "issues": "Есть отдушка",
        "status": False,
    }

    product = normalize_product_row(row, row_number=2)

    assert product["external_id"] == "product_001"
    assert product["name"] == "Мягкая пенка"
    assert product["brand"] == "Demo Brand"
    assert product["price_rub"] == 743
    assert product["budget_segment"] == "minimum"
    assert product["routine_step"] == "cleanser"
    assert product["product_url"] == "https://example.com/product"
    assert "Glycerin" in product["ingredients"]
    assert "contains_fragrance" in product["risk_flags"]
    assert product["source"] == "dataset"
    assert product["is_active"] is True


def test_product_dataset_mapper_accepts_russian_columns() -> None:
    row = {
        "Название": "Мягкая пенка",
        "Бренд": "Demo Brand",
        "Цена": "750 ₽",
        "Шаг ухода": "Очищение",
        "Типы кожи": "комбинированная; жирная",
        "Запросы": "жирный блеск; высыпания",
        "Состав": "glycerin; parfum",
        "Активен": "да",
    }

    product = normalize_product_row(row, row_number=2)

    assert product["name"] == "Мягкая пенка"
    assert product["brand"] == "Demo Brand"
    assert product["price_rub"] == 750
    assert product["routine_step"] == "cleanser"
    assert product["skin_types"] == ["combination", "oily"]
    assert product["concerns"] == ["sebum", "acne"]
    assert "contains_fragrance" in product["risk_flags"]
    assert product["source"] == "dataset"
    assert product["is_active"] is True


def test_product_dataset_mapper_generates_budget_from_price() -> None:
    row = {
        "name": "Basic SPF",
        "brand": "Demo Brand",
        "price": "2600",
        "routine_step": "spf",
    }

    product = normalize_product_row(row, row_number=3)

    assert product["budget_segment"] == "premium"
    assert product["routine_step"] == "spf"


def test_product_dataset_loader_reads_csv(tmp_path: Path) -> None:
    file_path = tmp_path / "products.csv"
    file_path.write_text(
        "Название;Бренд;Цена;Шаг ухода;Типы кожи;Запросы;Состав\n"
        "Пенка;Demo;700;Очищение;комбинированная;жирный блеск;glycerin\n",
        encoding="utf-8",
    )

    preview = load_products_from_dataset(file_path)

    assert len(preview.products) == 1
    assert preview.errors == []
    assert preview.products[0]["routine_step"] == "cleanser"


def test_product_dataset_loader_collects_invalid_rows(tmp_path: Path) -> None:
    file_path = tmp_path / "products.csv"
    file_path.write_text(
        "Название;Бренд;Цена;Шаг ухода\n"
        ";Demo;700;Очищение\n",
        encoding="utf-8",
    )

    preview = load_products_from_dataset(file_path)

    assert preview.products == []
    assert len(preview.errors) == 1
    assert preview.errors[0].reason == "Missing product name"


def test_product_dataset_loader_selects_products_sheet_not_consultations(tmp_path: Path) -> None:
    file_path = tmp_path / "products.xlsx"

    workbook = Workbook()
    consultation_sheet = workbook.active
    consultation_sheet.title = "Koyash"
    consultation_sheet.append(["consultation_id", "age", "skin_type", "concerns", "budget"])
    consultation_sheet.append(["cons_001", 22, "Жирная", "акне", "low"])

    products_sheet = workbook.create_sheet("Products")
    products_sheet.append(
        [
            "product_id",
            "name",
            "brand",
            "price_rub",
            "segment",
            "category",
            "link",
            "ingredients_raw",
            "main_actives",
            "functional_category",
        ]
    )
    products_sheet.append(
        [
            "product_001",
            "Мягкая пенка",
            "Demo",
            700,
            "бюджет",
            "Очищающий гель",
            "https://example.com",
            "Aqua, Glycerin",
            "Glycerin",
            "Мягкое очищение, гидратация",
        ]
    )

    workbook.save(file_path)

    preview = load_products_from_dataset(file_path)

    assert len(preview.products) == 1
    assert preview.errors == []
    assert preview.products[0]["external_id"] == "product_001"
    assert preview.products[0]["routine_step"] == "cleanser"
''')


print("Step 4.4 importer fixed for the real Products sheet.")