from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from scripts.refresh_product_metadata import (
    _column_map,
    _materialize_derived_fields,
    refresh,
)


def _worksheet():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        ["product_id", "price_rub", "segment", "volume_ml", "price_per_ml"]
    )
    return worksheet


def test_materializes_dataset_formulas_for_data_only_imports() -> None:
    worksheet = _worksheet()
    worksheet.append(["product_001", 743, '=IF(B2<1000,"бюджет","мидл")', 75, "=B2/D2"])

    changed = _materialize_derived_fields(
        worksheet,
        _column_map(worksheet),
        apply=True,
    )

    assert changed == 2
    assert worksheet["C2"].value == "бюджет"
    assert worksheet["E2"].value == 743 / 75


def test_derived_field_preview_does_not_mutate_cells() -> None:
    worksheet = _worksheet()
    worksheet.append(["minimum", 999, "stale", 100, 99])
    worksheet.append(["balanced", 1000, "stale", 100, 99])
    worksheet.append(["premium", 3001, "stale", 100, 99])

    changed = _materialize_derived_fields(
        worksheet,
        _column_map(worksheet),
        apply=False,
    )

    assert changed == 6
    assert [worksheet.cell(row=row, column=3).value for row in range(2, 5)] == [
        "stale",
        "stale",
        "stale",
    ]


def test_derived_only_apply_does_not_add_metadata_columns(tmp_path) -> None:
    worksheet = _worksheet()
    worksheet.append(["product_001", 743, "=formula", 75, "=formula"])
    dataset = tmp_path / "products.xlsx"
    worksheet.parent.save(dataset)

    report = refresh(
        SimpleNamespace(
            dataset=str(dataset),
            sheet="Sheet",
            apply=True,
            derived_only=True,
            only_id=None,
        )
    )

    saved = load_workbook(dataset, data_only=True)["Sheet"]
    headers = [cell.value for cell in saved[1]]
    assert headers == [
        "product_id",
        "price_rub",
        "segment",
        "volume_ml",
        "price_per_ml",
    ]
    assert saved["C2"].value == "бюджет"
    assert saved["E2"].value == 743 / 75
    assert report["summary"]["derived_cells_updated"] == 2
